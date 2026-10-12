from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(r"C:\Users\Admin\Downloads\Bao_cao_MCP_PyTutor_Tran_Khanh_v2.docx")
OUTPUT = Path(r"C:\Users\Admin\Downloads\Bao_cao_MCP_PyTutor_Tran_Khanh_v2_trinh_bay_thong_nhat.docx")
BLACK = "000000"
NAVY = "24476B"
PALE = "F2F5F9"
BORDER = "D9D9D9"


def text_inventory(doc):
    paragraphs = [p.text for p in doc.paragraphs]
    tables = [[ [cell.text for cell in row.cells] for row in table.rows ] for table in doc.tables]
    headers_footers = []
    for section in doc.sections:
        headers_footers.extend([
            [p.text for p in section.header.paragraphs],
            [p.text for p in section.footer.paragraphs],
            [p.text for p in section.first_page_header.paragraphs],
            [p.text for p in section.first_page_footer.paragraphs],
        ])
    return paragraphs, tables, headers_footers


def set_style_font(style, name, size, bold=None, color=BLACK, italic=None):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), name)


def set_run_font(run, name, size, color=BLACK):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), name)


def remove_paragraph_border(paragraph):
    ppr = paragraph._p.pPr
    if ppr is not None:
        borders = ppr.find(qn("w:pBdr"))
        if borders is not None:
            ppr.remove(borders)


def set_shading(element, fill):
    ppr = element.get_or_add_pPr() if element.tag == qn("w:p") else element.get_or_add_tcPr()
    shd = ppr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        ppr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tcpr = cell._tc.get_or_add_tcPr()
    margins = tcpr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tcpr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), BORDER)


def set_table_widths(table, available_width):
    old = [col.width.inches if col.width else 1.0 for col in table.columns]
    total = sum(old) or len(old)
    widths = [Inches(available_width * w / total) for w in old]
    table.autofit = False
    for col, width in zip(table.columns, widths):
        col.width = width
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = widths[min(idx, len(widths) - 1)]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def format_table(table, available_width):
    set_table_widths(table, available_width)
    set_table_borders(table)
    if table.rows:
        header_pr = table.rows[0]._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        header_pr.append(repeat)
    for row_index, row in enumerate(table.rows):
        trpr = row._tr.get_or_add_trPr()
        if trpr.find(qn("w:cantSplit")) is None:
            trpr.append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            set_cell_margins(cell)
            if row_index == 0:
                set_shading(cell._tc, NAVY)
            elif row_index % 2 == 0:
                set_shading(cell._tc, PALE)
            else:
                set_shading(cell._tc, "FFFFFF")
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    if any(ord(ch) >= 0x1F000 for ch in run.text):
                        set_run_font(run, "Segoe UI Emoji", 9)
                    else:
                        set_run_font(run, "Times New Roman", 9, "FFFFFF" if row_index == 0 else BLACK)
                    if row_index == 0:
                        run.bold = True


def is_code_or_diagram(text):
    stripped = text.strip()
    if not stripped:
        return False
    code_starts = (
        "from mcp.server", "mcp = ", "@mcp.tool()", "def run_code", "if __name__", "{\n", "{"
    )
    if stripped.startswith(code_starts):
        return True
    diagram_chars = "┌┐└┘│─▼"
    if any(ch in text for ch in diagram_chars):
        return True
    if "AI_A" in text and "Tool_" in text:
        return True
    return False


def format_heading(paragraph, level):
    style_name = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3"}[level]
    paragraph.style = style_name
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = paragraph.paragraph_format
    pf.space_before = Pt({1: 14, 2: 9, 3: 6}[level])
    pf.space_after = Pt({1: 6, 2: 4, 3: 3}[level])
    pf.line_spacing = 1.0
    pf.keep_with_next = True
    pf.keep_together = True
    for run in paragraph.runs:
        set_run_font(run, "Times New Roman", {1: 14, 2: 11.5, 3: 10.5}[level])
        run.bold = True


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    doc = Document(SOURCE)
    original = text_inventory(doc)
    section = doc.sections[0]
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)

    set_style_font(doc.styles["Normal"], "Times New Roman", 10)
    normal_pf = doc.styles["Normal"].paragraph_format
    normal_pf.space_before = Pt(0)
    normal_pf.space_after = Pt(4)
    normal_pf.line_spacing = 1.08
    normal_pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    set_style_font(doc.styles["Title"], "Times New Roman", 16, bold=True)
    title_pf = doc.styles["Title"].paragraph_format
    title_pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_pf.space_before = Pt(0)
    title_pf.space_after = Pt(8)
    title_pf.keep_with_next = True

    set_style_font(doc.styles["Subtitle"], "Times New Roman", 11, color=BLACK)
    subtitle_pf = doc.styles["Subtitle"].paragraph_format
    subtitle_pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_pf.space_before = Pt(0)
    subtitle_pf.space_after = Pt(6)
    subtitle_pf.keep_with_next = True

    for name, size in (("Heading 1", 14), ("Heading 2", 11.5), ("Heading 3", 10.5)):
        set_style_font(doc.styles[name], "Times New Roman", size, bold=True)
        pf = doc.styles[name].paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.space_before = Pt({"Heading 1": 14, "Heading 2": 9, "Heading 3": 6}[name])
        pf.space_after = Pt({"Heading 1": 6, "Heading 2": 4, "Heading 3": 3}[name])
        pf.keep_with_next = True
        pf.keep_together = True
        pf.line_spacing = 1.0

    if "Code Block" not in [style.name for style in doc.styles]:
        code_style = doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code_style = doc.styles["Code Block"]
    set_style_font(code_style, "Consolas", 8.5)
    code_pf = code_style.paragraph_format
    code_pf.left_indent = Inches(0.12)
    code_pf.right_indent = Inches(0.08)
    code_pf.space_before = Pt(3)
    code_pf.space_after = Pt(4)
    code_pf.line_spacing = 1.0
    code_pf.keep_together = True

    available_width = section.page_width.inches - section.left_margin.inches - section.right_margin.inches
    for table in doc.tables:
        format_table(table, available_width)

    in_questions = False
    for idx, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue
        remove_paragraph_border(paragraph)

        if idx == 0:
            paragraph.style = "Title"
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_run_font(run, "Times New Roman", 16)
                run.bold = True
            continue

        if idx == 1:
            paragraph.style = "Subtitle"
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_run_font(run, "Times New Roman", 11)
            continue

        if text.startswith("PHẦN TRAO ĐỔI VỚI CÁC CÂU HỎI"):
            in_questions = True
            format_heading(paragraph, 1)
            continue
        if text.startswith("PHẦN II:"):
            in_questions = False
            format_heading(paragraph, 1)
            continue
        if text.startswith("Bước ") and re.match(r"^Bước\s+\d+:", text):
            format_heading(paragraph, 3)
            continue
        if re.match(r"^\d+\.\d+(?:\.\d+)?\s", text):
            format_heading(paragraph, 2)
            continue
        if re.match(r"^\d+\.\s", text):
            format_heading(paragraph, 2 if in_questions else 1)
            continue
        if text.startswith(("So sánh:", "Tài liệu tham khảo chính:")):
            format_heading(paragraph, 2)
            continue

        if is_code_or_diagram(paragraph.text):
            paragraph.style = code_style
            for run in paragraph.runs:
                set_run_font(run, "Consolas", 8.5)
            if any(ch in paragraph.text for ch in "┌┐└┘│─▼") or "AI_A" in paragraph.text:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            continue

        if paragraph.style.name == "List Bullet":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf = paragraph.paragraph_format
            pf.left_indent = Inches(0.28)
            pf.first_line_indent = Inches(-0.18)
            pf.space_before = Pt(0)
            pf.space_after = Pt(2)
            pf.line_spacing = 1.05
        else:
            paragraph.style = "Normal"
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT
                if "\n" in paragraph.text or "\t" in paragraph.text
                else WD_ALIGN_PARAGRAPH.JUSTIFY
            )
            pf = paragraph.paragraph_format
            pf.space_before = Pt(0)
            pf.space_after = Pt(4)
            pf.line_spacing = 1.08
        for run in paragraph.runs:
            if any(ord(ch) >= 0x1F000 for ch in run.text):
                set_run_font(run, "Segoe UI Emoji", 10)
            else:
                set_run_font(run, "Times New Roman", 10)

    doc.save(OUTPUT)

    result = Document(OUTPUT)
    if text_inventory(result) != original:
        raise AssertionError("Text in paragraphs, tables, headers, or footers changed during formatting.")
    if len(result.inline_shapes) != len(doc.inline_shapes):
        raise AssertionError("The number of images changed during formatting.")
    print(f"Saved: {OUTPUT}")
    print(f"Paragraphs: {len(result.paragraphs)} | Tables: {len(result.tables)} | Images: {len(result.inline_shapes)}")
    print("Text check: all paragraph, table, header, and footer text is unchanged.")


if __name__ == "__main__":
    main()
