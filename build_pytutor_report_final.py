from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn


ROOT = Path(r"C:\Users\Admin\Downloads\Sales-support-AI-pilot")
TEMPLATE = Path(r"C:\Users\Admin\Downloads\Bao_cao_Tran_Khanh.docx")
WIRE = ROOT / "pytutor-main-wireframe.png"
OUTPUT = ROOT / "Bao_cao_PyTutor_Tran_Khanh_hoan_chinh.docx"


def add_body(doc, text, style="Normal"):
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p


def add_heading(doc, text, level=1):
    style = "heading1" if level == 1 else "heading2"
    return doc.add_paragraph(text, style=style)


def add_labeled_paragraph(doc, label, text):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.keep_together = True
    r = p.add_run(label)
    r.bold = True
    p.add_run(text)
    return p


def add_reference(doc, text):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.22)
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)
    return p


def main():
    if not TEMPLATE.exists() or not WIRE.exists():
        raise FileNotFoundError("The retained report template or PyTutor wireframe is missing.")

    doc = Document(TEMPLATE)
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)

    doc.core_properties.title = (
        "Thiết kế và đánh giá hệ thống gia sư thông minh dựa trên Model Context Protocol MCP "
        "hỗ trợ học lập trình Python"
    )
    doc.core_properties.author = "Trần Khánh"
    doc.core_properties.subject = "Báo cáo thiết kế PyTutor và các bản demo MCP"

    try:
        title_style = doc.styles["Title"]
    except KeyError:
        title_style = doc.styles.add_style("Title", WD_STYLE_TYPE.PARAGRAPH)
    title_style.font.name = "Times New Roman"
    title_style.font.size = Pt(13)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title = doc.add_paragraph(style=title_style)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    run = title.add_run(
        "THIẾT KẾ VÀ ĐÁNH GIÁ HỆ THỐNG GIA SƯ THÔNG MINH DỰA TRÊN "
        "MODEL CONTEXT PROTOCOL MCP HỖ TRỢ HỌC LẬP TRÌNH PYTHON"
    )
    run.bold = True
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0, 0, 0)

    author = doc.add_paragraph(style="author")
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.add_run("Trần Khánh")

    abstract = doc.add_paragraph(style="abstract")
    abstract.paragraph_format.keep_together = True
    abstract.add_run("Tóm tắt. ").bold = True
    abstract.add_run(
        "Báo cáo trình bày thiết kế PyTutor, một hệ thống gia sư thông minh trực tuyến hỗ trợ sinh viên "
        "luyện lập trình Python bằng bài tập, kiểm thử và gợi ý theo từng bước. Model Context Protocol "
        "(MCP) được đề xuất làm lớp giao tiếp giữa ứng dụng gia sư và các công cụ như chạy mã, kiểm thử "
        "và phân tích cú pháp; các quyết định sư phạm thuộc về mô hình người học, quy tắc trợ giúp và "
        "mô hình ngôn ngữ. Báo cáo phân biệt hai nguyên mẫu hiện có: một demo trao đổi MCP thực qua "
        "stdio với công cụ run_code, và một demo luồng gia sư dùng các hàm công cụ mô phỏng để tạo "
        "kết quả kiểm thử, phân tích AST và gọi dịch vụ Qwen. Vì hai phần chưa được tích hợp thành một "
        "hệ thống hoàn chỉnh, kết quả hiện tại chỉ chứng minh tính khả thi ở mức giao thức và luồng "
        "minh họa; chưa chứng minh hiệu quả học tập hoặc độ an toàn khi chạy mã không tin cậy."
    )

    keywords = doc.add_paragraph(style="keywords")
    keywords.add_run("Từ khóa: ").bold = True
    keywords.add_run("hệ thống gia sư thông minh; MCP; học lập trình; Python; gợi ý thích ứng")

    add_heading(doc, "Giới thiệu")
    add_body(
        doc,
        "Người mới học lập trình thường nhận được thông báo lỗi hoặc kết quả sai mà chưa biết cách "
        "lần theo nguyên nhân. Một chatbot chỉ nhìn đoạn mã có thể giải thích cú pháp, nhưng không "
        "nhất thiết biết chương trình đã chạy ra sao trên các trường hợp kiểm thử. PyTutor hướng đến "
        "một không gian làm bài trong đó sinh viên có thể đọc đề, viết mã, chạy thử, xem kết quả và "
        "yêu cầu gợi ý mà không phải chuyển ngữ cảnh sang một công cụ khác."
    )
    add_body(
        doc,
        "Mục tiêu của đề tài là thiết kế một gia sư lập trình Python có thể kết hợp dữ liệu bài tập, "
        "mã nguồn, kết quả thực thi và lịch sử làm bài để hỗ trợ quá trình sửa lỗi. MCP được khảo sát "
        "và thử nghiệm như giao thức tích hợp giữa ứng dụng chủ và các công cụ. Phạm vi hiện tại gồm "
        "wireframe chức năng chính, nguyên mẫu giao thức cục bộ và nguyên mẫu minh họa luồng tạo hint. "
        "Báo cáo không xem các nguyên mẫu này là một sản phẩm đã triển khai đầy đủ hoặc một nghiên cứu "
        "đã đo lường kết quả học tập."
    )

    add_heading(doc, "Cơ sở và công nghệ liên quan")
    add_heading(doc, "Hệ thống gia sư thông minh", level=2)
    add_body(
        doc,
        "Một ITS cần phản hồi theo tiến trình giải bài và có thể cập nhật hiểu biết về trạng thái "
        "người học. Mô hình tutoring của VanLehn mô tả vòng tương tác trong lúc giải từng bước, nơi "
        "hệ thống có thể đưa phản hồi, gợi ý và cập nhật Student Model [4]. Các ITS dạy lập trình "
        "khác nhau về phạm vi kiến thức, hình thức phản hồi và tính năng hỗ trợ; vì vậy việc có một "
        "LLM hoặc compiler riêng lẻ chưa đủ để khẳng định một ứng dụng là gia sư thích ứng [5]."
    )
    add_heading(doc, "Vai trò của MCP", level=2)
    add_body(
        doc,
        "MCP định nghĩa mô hình Host-Client-Server dựa trên JSON-RPC. Host là ứng dụng điều phối; "
        "Host tạo MCP Client để kết nối tới Server; Server cung cấp các khả năng đã khai báo. Các "
        "primitive chính gồm Tools để gọi hành động, Resources để đọc dữ liệu và Prompts để cung cấp "
        "mẫu hướng dẫn [1], [2]. Client có thể liệt kê Tools bằng tools/list rồi gọi một Tool bằng "
        "tools/call. Giao thức không tự chạy chương trình, tạo dữ liệu đáng tin cậy hay quyết định "
        "cách dạy; các việc đó thuộc về phần mềm thực thi và logic của ứng dụng."
    )
    add_body(
        doc,
        "MCP hỗ trợ vận chuyển qua stdio cho tiến trình cục bộ và Streamable HTTP cho kết nối từ xa "
        "[3]. Nguyên mẫu trong dự án dùng stdio. Nếu chuyển sang dịch vụ web, cần triển khai và bảo "
        "vệ endpoint từ xa, bao gồm xác thực và kiểm soát nguồn truy cập. MCP có thể chuẩn hóa cách "
        "Host kết nối nhiều công cụ, nhưng bản thân giao thức không đảm bảo hệ thống sẽ chính xác "
        "hơn hoặc dễ mở rộng hơn trong mọi thiết kế."
    )

    add_heading(doc, "Thiết kế hệ thống PyTutor")
    add_heading(doc, "Người dùng và chức năng chính", level=2)
    add_body(
        doc,
        "Đối tượng sử dụng chính là sinh viên đang học Python. Trong một bài tập, sinh viên đọc yêu "
        "cầu và ví dụ, chỉnh sửa mã trong trình soạn thảo, chọn Chạy thử để xem output hoặc chọn Nộp "
        "bài để đối chiếu với bộ kiểm thử. Khi chưa tìm được nguyên nhân, sinh viên yêu cầu gợi ý; "
        "hệ thống nên chỉ đưa thông tin vừa đủ để người học tiếp tục tự sửa mã."
    )
    add_heading(doc, "Kiến trúc đề xuất", level=2)
    add_body(
        doc,
        "Giao diện trình duyệt gửi thao tác tới backend PyTutor, đóng vai trò MCP Host. Host quản lý "
        "phiên làm bài, lấy ngữ cảnh cần thiết, áp dụng quy tắc sư phạm và gọi dịch vụ mô hình ngôn "
        "ngữ. MCP Client trong Host kết nối tới MCP Server. Server có thể cung cấp các Tool như "
        "run_tests, analyze_ast và các Resource như nội dung đề bài hoặc lịch sử nộp bài. Phần mềm "
        "kiểm thử, cơ sở dữ liệu và môi trường chạy mã là các thành phần riêng phía sau Server; MCP "
        "chỉ quy định cách chúng được cung cấp qua giao thức."
    )
    add_body(
        doc,
        "Khi sinh viên yêu cầu hint, Host ghép đề bài, mã hiện tại, kết quả kiểm thử, lỗi phân tích "
        "và các hint trước đó nếu có. Student Model lưu thông tin tiến trình; Pedagogical Model "
        "chọn mức trợ giúp; mô hình ngôn ngữ diễn đạt phản hồi. Các thành phần Student Model, "
        "Pedagogical Model, cơ sở dữ liệu và tích hợp MCP đầy đủ thuộc kiến trúc mục tiêu, chưa được "
        "chứng minh là đã hoàn thiện trong các demo hiện có."
    )
    add_heading(doc, "Wireframe chức năng chính", level=2)
    add_body(
        doc,
        "Wireframe minh họa một màn hình gồm đề bài, trình soạn thảo và kết quả kiểm thử, cùng khu "
        "vực AI Tutor. Người học có thể yêu cầu gợi ý hoặc yêu cầu giải thích rõ hơn. Đây là bản phác "
        "thảo giao diện, không phải ảnh chụp một hệ thống đã triển khai."
    )
    picture = doc.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.keep_together = True
    picture.add_run().add_picture(str(WIRE), width=Inches(4.7))
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_together = True
    caption.paragraph_format.space_after = Pt(8)
    cap = caption.add_run("Hình 1. Wireframe chức năng làm bài và nhận gợi ý thích ứng")
    cap.italic = True

    add_heading(doc, "Luồng yêu cầu gợi ý", level=2)
    add_labeled_paragraph(
        doc,
        "Bước 1 - Thu thập ngữ cảnh. ",
        "Host nhận mã hiện tại, đề bài, kết quả chạy hoặc kiểm thử và trạng thái bài làm."
    )
    add_labeled_paragraph(
        doc,
        "Bước 2 - Gọi công cụ. ",
        "MCP Client gọi các Tool phù hợp; Server trả kết quả có cấu trúc để Host kiểm tra."
    )
    add_labeled_paragraph(
        doc,
        "Bước 3 - Chọn cách trợ giúp. ",
        "Pedagogical Model chọn mức gợi ý, sau đó LLM tạo phản hồi ngắn dựa trên bằng chứng."
    )
    add_labeled_paragraph(
        doc,
        "Bước 4 - Người học tiếp tục. ",
        "Sinh viên xem hint, sửa mã và chạy lại; kết quả được ghi nhận để dùng trong vòng hỗ trợ kế tiếp."
    )
    add_body(
        doc,
        "Wireframe hiện thể hiện thanh gợi ý bốn mức, trong khi demo luồng gia sư dùng ba cấu hình "
        "prompt. Trước khi tích hợp giao diện và backend, cần thống nhất số mức và nội dung của từng "
        "mức. Một cách chia phù hợp là: gợi mở kiến thức; chỉ vùng cần xem xét; giải thích nguyên nhân; "
        "đưa ra các bước sửa hoặc mã giả, không cung cấp ngay lời giải hoàn chỉnh. Đây là đề xuất "
        "thiết kế, chưa phải kết quả đánh giá với người học."
    )

    add_heading(doc, "Nguyên mẫu và mức độ triển khai hiện tại")
    add_heading(doc, "Demo giao thức MCP thực", level=2)
    add_body(
        doc,
        "Trong mcp_visual_demo.py, Streamlit đóng vai trò giao diện demo và MCP Client dùng Python "
        "SDK kết nối qua stdio tới mcp_server_pytutor.py. Server khai báo một Tool là run_code. Client "
        "khởi tạo phiên, liệt kê Tool và gọi run_code; phần log trong giao diện được dựng lại từ kết "
        "quả SDK để minh họa các bước, không phải bản ghi thô đọc trực tiếp từ kênh stdio. Đây là "
        "phần hiện thực hóa MCP thực trong thư mục dự án. Tool "
        "run_code thực thi Python bằng exec trong tiến trình Server; việc chạy ở tiến trình con "
        "không đồng nghĩa với sandbox bảo mật. Demo này chưa kết nối LLM, cơ sở dữ liệu hay Student "
        "Model."
    )
    add_heading(doc, "Demo luồng gia sư", level=2)
    add_body(
        doc,
        "Trong pytutor_mcp_demo.py, giao diện cho thấy bài tập tính tổng, các ca kiểm thử, phân tích "
        "AST và các mức yêu cầu gợi ý. Hàm Python cục bộ thực hiện run_tests và analyze_ast; chương "
        "trình tự tạo các đối tượng JSON để trình bày log giống thông điệp MCP, nhưng nhánh này "
        "không dùng MCP Client/Server của SDK để trao đổi các Tool đó. Ứng dụng có đoạn gọi một API "
        "Qwen bên ngoài. Do đó, đây là demo luồng tương tác và prompt, không phải bằng chứng rằng "
        "các Tool tutoring đã được tích hợp qua MCP."
    )
    add_heading(doc, "Kết quả và thảo luận", level=1)
    add_body(
        doc,
        "Hai nguyên mẫu chứng minh được hai việc riêng: demo giao thức cho thấy Client có thể kết "
        "nối tới Server cục bộ, khám phá Tool và nhận kết quả gọi Tool; demo gia sư cho thấy có thể "
        "đưa kết quả kiểm thử và nhận xét AST vào prompt để tạo phản hồi theo mức. Chưa có một luồng "
        "duy nhất dùng MCP thực từ giao diện sinh viên tới các Tool kiểm thử, phân tích và dịch vụ "
        "LLM. Vì vậy, kết luận phù hợp hiện nay là tính khả thi kỹ thuật ở mức nguyên mẫu, không phải "
        "mức hoàn thiện của sản phẩm."
    )
    add_body(
        doc,
        "Bản v2 ghi nhận thời gian xấp xỉ 1 ms khi chạy trực tiếp và khoảng 300-500 ms qua demo MCP. "
        "Các con số này phụ thuộc máy thử nghiệm, thời gian khởi động tiến trình con và cách đo; tài "
        "liệu chưa nêu đủ số lần lặp và điều kiện đo để xem đây là benchmark. Chúng cũng không so "
        "sánh công bằng hai hệ thống tương đương: MCP bổ sung khâu giao thức và khởi động, còn chạy "
        "trực tiếp bỏ qua các khâu đó. Không thể suy ra từ số liệu này rằng MCP luôn chậm như vậy "
        "hoặc một server chạy liên tục sẽ có độ trễ cụ thể nào."
    )
    add_heading(doc, "Hạn chế và kế hoạch đánh giá", level=1)
    add_body(
        doc,
        "Hạn chế lớn nhất là hai demo chưa được nối thành một kiến trúc MCP thống nhất. Các mã chạy "
        "bằng exec không bị cô lập khỏi hệ điều hành bằng container hay cơ chế sandbox; vì thế không "
        "nên dùng chúng để chạy mã sinh viên trên máy chủ dùng chung. Trước khi thử nghiệm với người "
        "dùng, cần cô lập tiến trình và áp dụng giới hạn thời gian, CPU, bộ nhớ, tệp và mạng. Cần bổ "
        "sung kiểm soát đầu vào, xác thực, quyền truy cập và ghi nhật ký phù hợp."
    )
    add_body(
        doc,
        "Giai đoạn kỹ thuật tiếp theo là đưa run_tests và analyze_ast vào MCP Server, kết nối bằng "
        "MCP Client thật, xác nhận schema đầu vào/đầu ra và kiểm tra các ca thành công, lỗi cú pháp, "
        "timeout cùng lời gọi Tool không hợp lệ. Sau đó mới tích hợp Student Model, lưu lịch sử và "
        "quy tắc chọn hint. Khi triển khai qua mạng, có thể chuyển sang Streamable HTTP với xác thực "
        "và các kiểm soát bảo mật của transport [3]."
    )
    add_body(
        doc,
        "Đánh giá nên tách thành ba lớp: (1) kiểm tra đúng giao thức và độ tin cậy của lời gọi Tool; "
        "(2) đánh giá chất lượng hint bằng rubric, gồm tính chính xác, mức độ gợi mở và việc tránh "
        "lộ lời giải; (3) nghiên cứu người học để xem hệ thống có hỗ trợ hiểu bài và tự sửa lỗi hay "
        "không. Chỉ lớp thứ ba mới cung cấp bằng chứng về hiệu quả học tập; cần thiết kế hoạt động, "
        "tiêu chí và quy trình thu thập dữ liệu trước khi đưa ra kết luận."
    )

    add_heading(doc, "Kết luận")
    add_body(
        doc,
        "PyTutor đề xuất kết hợp giao diện làm bài Python, kiểm tra tự động và gợi ý thích ứng; MCP "
        "đảm nhiệm lớp giao tiếp để Host sử dụng dữ liệu và công cụ. Mã nguồn dự án hiện có một demo "
        "MCP thực cho Tool run_code và một demo gia sư riêng sử dụng các hàm kiểm thử, phân tích AST "
        "cục bộ cùng lời gọi Qwen. Việc phân biệt rõ hai nguyên mẫu giúp xác định bước hoàn thiện "
        "tiếp theo: tích hợp Tool tutoring qua MCP, xây môi trường thực thi an toàn, thống nhất mức "
        "gợi ý và đánh giá hệ thống bằng kiểm thử kỹ thuật lẫn nghiên cứu với người học."
    )

    add_heading(doc, "Tài liệu tham khảo")
    add_reference(doc, "[1] Model Context Protocol, “Architecture,” MCP Specification, 2025-06-18. https://modelcontextprotocol.io/specification/2025-06-18/architecture")
    add_reference(doc, "[2] Model Context Protocol, “Tools,” MCP Specification, 2025-06-18. https://modelcontextprotocol.io/specification/2025-06-18/server/tools")
    add_reference(doc, "[3] Model Context Protocol, “Transports,” MCP Specification, 2025-06-18. https://modelcontextprotocol.io/specification/2025-06-18/basic/transports")
    add_reference(doc, "[4] K. VanLehn, “The Behavior of Tutoring Systems,” International Journal of Artificial Intelligence in Education, vol. 16, no. 3, pp. 227-265, 2006. doi:10.3233/IRG-2006-16(3)02.")
    add_reference(doc, "[5] T. Crow, A. Luxton-Reilly, and B. Wünsche, “Intelligent Tutoring Systems for Programming Education: A Systematic Review,” in Proceedings of the 20th Australasian Computing Education Conference, pp. 53-62, 2018. doi:10.1145/3160489.3160492.")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
