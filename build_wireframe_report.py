from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r"C:\Users\Admin\Downloads\Sales-support-AI-pilot\Bao-cao-wireframe-chuc-nang-chinh-PyTutor.docx"
IMAGE = r"C:\Users\Admin\Downloads\Sales-support-AI-pilot\pytutor-main-wireframe.png"

BLUE = "2E5F88"
DARK = "1F2933"
LIGHT = "EAF1F7"
GRAY = "66717C"

doc = Document()
section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(2.2)
section.right_margin = Cm(2.0)
section.header_distance = Cm(1.0)
section.footer_distance = Cm(1.0)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(DARK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15

for name, size, before, after in [
    ("Heading 1", 16, 14, 7),
    ("Heading 2", 13, 10, 5),
]:
    st = styles[name]
    st.font.name = "Arial"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(BLUE)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def set_table_geometry(table, widths_cm):
    table.autofit = False
    total_dxa = sum(int(w / 2.54 * 1440) for w in widths_cm)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for w in widths_cm:
        dxa = int(w / 2.54 * 1440)
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(dxa))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            dxa = int(widths_cm[i] / 2.54 * 1440)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(dxa))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_field(paragraph, instr):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instr
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, text, separate, end])

# Quiet running header/footer.
hp = section.header.paragraphs[0]
hp.text = "PYTUTOR MCP | BÁO CÁO WIREFRAME"
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for r in hp.runs:
    r.font.name = "Arial"; r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string(GRAY)
fp = section.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run("Trang ")
fr.font.name = "Arial"; fr.font.size = Pt(9); fr.font.color.rgb = RGBColor.from_string(GRAY)
add_field(fp, "PAGE")

# Title block.
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after = Pt(4)
r = p.add_run("BÁO CÁO WIREFRAME CHỨC NĂNG CHÍNH")
r.font.name = "Arial"; r.font.size = Pt(22); r.font.bold = True; r.font.color.rgb = RGBColor.from_string(DARK)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(14)
r = p.add_run("Hệ thống gia sư thông minh hỗ trợ học lập trình Python dựa trên Model Context Protocol (MCP)")
r.font.name = "Arial"; r.font.size = Pt(12.5); r.font.italic = True; r.font.color.rgb = RGBColor.from_string(GRAY)

meta = doc.add_table(rows=3, cols=2)
meta.style = "Table Grid"
meta.cell(0,0).text = "Sinh viên thực hiện"
meta.cell(0,1).text = "........................................................"
meta.cell(1,0).text = "Môn học / Lớp"
meta.cell(1,1).text = "........................................................"
meta.cell(2,0).text = "Ngày thực hiện"
meta.cell(2,1).text = "13/08/2026"
set_table_geometry(meta, [4.2, 12.6])
for row in meta.rows:
    set_cell_shading(row.cells[0], LIGHT)
    row.cells[0].paragraphs[0].runs[0].bold = True

doc.add_heading("1. Mục đích", level=1)
doc.add_paragraph(
    "Tài liệu trình bày bản thiết kế giao diện sơ bộ cho chức năng chính của hệ thống: sinh viên làm bài lập trình Python và nhận gợi ý thích ứng từ gia sư AI trên cùng một màn hình. Wireframe được xây dựng theo phong cách low-fidelity của Balsamiq nhằm làm rõ bố cục, thành phần và luồng tương tác trước khi phát triển phần mềm."
)

doc.add_heading("2. Chức năng chính được đề xuất", level=1)
lead = doc.add_paragraph()
lead.paragraph_format.left_indent = Cm(0.35)
lead.paragraph_format.right_indent = Cm(0.35)
lead.paragraph_format.space_before = Pt(3)
lead.paragraph_format.space_after = Pt(9)
lead_run = lead.add_run("Sinh viên đọc đề bài, viết mã, chạy thử hoặc nộp bài, xem kết quả kiểm thử và yêu cầu gợi ý thích ứng để tự sửa lỗi.")
lead_run.bold = True
lead_run.font.color.rgb = RGBColor.from_string(BLUE)

doc.add_heading("3. Luồng sử dụng", level=1)
flow = doc.add_paragraph()
flow.alignment = WD_ALIGN_PARAGRAPH.CENTER
flow.paragraph_format.space_before = Pt(5)
flow.paragraph_format.space_after = Pt(9)
rr = flow.add_run("Đọc đề  →  Viết mã  →  Chạy/Nộp bài  →  Xem kiểm thử  →  Nhận gợi ý  →  Sửa mã")
rr.bold = True; rr.font.size = Pt(11.5)

doc.add_page_break()
doc.add_heading("4. Wireframe giao diện chức năng chính", level=1)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(4)
p.add_run().add_picture(IMAGE, width=Cm(16.7))
cap = doc.add_paragraph("Hình 1. Wireframe màn hình làm bài và nhận gợi ý thích ứng")
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap.paragraph_format.space_after = Pt(9)
for r in cap.runs:
    r.font.size = Pt(9); r.font.italic = True; r.font.color.rgb = RGBColor.from_string(GRAY)

doc.add_heading("5. Mô tả các khu vực giao diện", level=1)
rows = [
    ("1. Đề bài", "Hiển thị nội dung bài tập, yêu cầu, ví dụ, khái niệm liên quan và mục tiêu học tập."),
    ("2. Trình soạn thảo", "Cho phép sinh viên viết mã Python, chạy thử và nộp bài."),
    ("3. Kết quả kiểm thử", "Thông báo số kiểm thử đạt/chưa đạt và dữ liệu đầu vào gây lỗi."),
    ("4. AI Tutor", "Tạo gợi ý theo mức độ dựa trên mã hiện tại, kết quả kiểm thử và lịch sử làm bài; ưu tiên định hướng thay vì cung cấp đáp án ngay."),
]
tbl = doc.add_table(rows=1, cols=2)
tbl.style = "Table Grid"
tbl.rows[0].cells[0].text = "Khu vực"
tbl.rows[0].cells[1].text = "Vai trò"
for a, b in rows:
    cells = tbl.add_row().cells
    cells[0].text = a
    cells[1].text = b
set_table_geometry(tbl, [4.0, 12.8])
for cell in tbl.rows[0].cells:
    set_cell_shading(cell, BLUE)
    for run in cell.paragraphs[0].runs:
        run.font.color.rgb = RGBColor(255,255,255); run.bold = True
for row in tbl.rows[1:]:
    row.cells[0].paragraphs[0].runs[0].bold = True

doc.add_heading("6. Kết luận", level=1)
doc.add_paragraph(
    "Wireframe tập trung đúng chức năng cốt lõi của đề tài và thể hiện đầy đủ vòng lặp học tập: thực hành, nhận phản hồi, nhận gợi ý và chỉnh sửa. Các màn hình phụ như đăng nhập, quản lý tài khoản và dashboard giảng viên chưa được đưa vào báo cáo này vì không thuộc phạm vi yêu cầu hiện tại."
)

doc.core_properties.title = "Báo cáo wireframe chức năng chính - PyTutor MCP"
doc.core_properties.subject = "Đồ án hệ thống gia sư thông minh hỗ trợ học lập trình Python"
doc.core_properties.author = "Sinh viên thực hiện"
doc.save(OUT)
print(OUT)
