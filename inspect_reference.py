from docx import Document
from docx.oxml.ns import qn
from pathlib import Path
import json, zipfile

p = Path(r"C:\Users\Admin\Downloads\Bao_cao_Tran_Khanh.docx")
d = Document(p)
data = {
  "paragraphs": [], "tables": [], "sections": [], "styles_used": {},
  "headers": [], "footers": []
}
for i, para in enumerate(d.paragraphs):
    if para.text.strip():
        fmt = para.paragraph_format
        runs = []
        for r in para.runs:
            if r.text:
                runs.append({"text": r.text, "font": r.font.name, "size": r.font.size.pt if r.font.size else None, "bold": r.bold, "italic": r.italic})
        data["paragraphs"].append({"i": i, "text": para.text, "style": para.style.name, "align": str(para.alignment), "before": fmt.space_before.pt if fmt.space_before else None, "after": fmt.space_after.pt if fmt.space_after else None, "line": str(fmt.line_spacing), "runs": runs})
        data["styles_used"][para.style.name] = data["styles_used"].get(para.style.name, 0) + 1
for ti, table in enumerate(d.tables):
    data["tables"].append([[cell.text for cell in row.cells] for row in table.rows])
for si, s in enumerate(d.sections):
    data["sections"].append({"width": s.page_width.inches, "height": s.page_height.inches, "top": s.top_margin.inches, "bottom": s.bottom_margin.inches, "left": s.left_margin.inches, "right": s.right_margin.inches, "different_first": s.different_first_page_header_footer})
    data["headers"].append([x.text for x in s.header.paragraphs])
    data["footers"].append([x.text for x in s.footer.paragraphs])
Path(r"C:\Users\Admin\Downloads\Sales-support-AI-pilot\reference-inspection.json").write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(data, ensure_ascii=False, indent=2))
