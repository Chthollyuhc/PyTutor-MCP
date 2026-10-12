from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path(r"C:\Users\Admin\Downloads\Sales-support-AI-pilot\Bao_cao_wireframe_PyTutor_Tran_Khanh.docx")
IMAGE = Path(r"C:\Users\Admin\Downloads\Sales-support-AI-pilot\pytutor-main-wireframe.png")

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(4.4)
sec.top_margin = sec.bottom_margin = Cm(3.5)
sec.header_distance = Cm(1.25)
sec.footer_distance = Cm(1.25)
sec.different_first_page_header_footer = True

def font(run, size=10, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size); run.bold = bold; run.italic = italic

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(2)
normal.paragraph_format.line_spacing = 1.0

def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t"); txt.text = "1"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, sep, txt, end]); font(run, 9)

hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
add_page_field(hp)

def para(text="", align=None, before=0, after=2, indent=0, first=0):
    p = doc.add_paragraph()
    if align is not None: p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.left_indent = Cm(indent)
    p.paragraph_format.first_line_indent = Cm(first)
    if text: font(p.add_run(text))
    return p

def heading(number, title, level=1):
    p = para(before=10 if level == 1 else 7, after=5)
    p.paragraph_format.keep_with_next = True
    if level == 1:
        font(p.add_run(f"{number}     {title}"), 12, True)
    else:
        font(p.add_run(f"{number}     {title}"), 10.5, True)
    return p

# Title block modeled on the reference paper.
p = para(align=WD_ALIGN_PARAGRAPH.CENTER, before=52, after=9)
font(p.add_run("Design of the Main Interface for an MCP-Based Intelligent Tutoring System for Python Programming Education"), 14, True)
p = para(align=WD_ALIGN_PARAGRAPH.CENTER, after=7)
font(p.add_run("Tran Khanh"), 10)
p = para(align=WD_ALIGN_PARAGRAPH.CENTER, after=1)
font(p.add_run("Faculty of Information Technology, College of Technology,"), 9.5)
p = para(align=WD_ALIGN_PARAGRAPH.CENTER, after=16)
font(p.add_run("National Economics University, Vietnam"), 9.5)

p = para(after=8)
font(p.add_run("Abstract. "), 9.5, True)
font(p.add_run("Báo cáo trình bày wireframe cho chức năng chính của hệ thống gia sư thông minh hỗ trợ học lập trình Python. Trên cùng một giao diện, sinh viên có thể đọc đề, viết mã, chạy thử, nộp bài, xem kết quả kiểm thử và nhận gợi ý thích ứng. Bản thiết kế tập trung vào vòng lặp học tập cốt lõi, không mở rộng sang các màn hình phụ."), 9.5)
p = para(after=13)
font(p.add_run("Keywords: "), 9.5, True)
font(p.add_run("Intelligent Tutoring System; Model Context Protocol; Python Programming; Adaptive Hints; Wireframe"), 9.5)

heading("1", "Introduction")
para("Trong quá trình học lập trình, thông báo lỗi hoặc kết quả kiểm thử thường chỉ cho sinh viên biết chương trình chưa đúng mà chưa giúp họ xác định nguyên nhân. Hệ thống được đề xuất kết hợp môi trường thực hành Python với gia sư AI để cung cấp phản hồi theo từng bước và theo dõi quá trình sửa bài của người học.", after=4)
para("Wireframe nhằm thể hiện giao diện mà sinh viên trực tiếp sử dụng, các thao tác chính và phản hồi tương ứng của hệ thống trước khi tiến hành xây dựng sản phẩm.")

heading("2", "Main Function")
para("Chức năng chính là làm bài lập trình và nhận gợi ý thích ứng. Sinh viên đọc yêu cầu bài tập, viết mã Python trong trình soạn thảo, chạy thử hoặc nộp bài để kiểm tra, sau đó yêu cầu AI Tutor hỗ trợ khi gặp lỗi.")

heading("2.1", "User Actions", 2)
actions = [
    ("Run Code", "chạy mã trong môi trường cách ly và trả lại đầu ra hoặc lỗi thực thi."),
    ("Submit", "chạy bộ kiểm thử, hiển thị test đạt/chưa đạt và lưu kết quả lần làm bài."),
    ("Request Hint", "thu thập mã hiện tại, kết quả kiểm thử và lịch sử thử để tạo gợi ý phù hợp."),
]
for label, text in actions:
    p = para(indent=0.55, first=-0.35, after=3)
    font(p.add_run("•  "), 10)
    font(p.add_run(f"{label}: "), 10, True)
    font(p.add_run(text), 10)

heading("2.2", "Adaptive Hint", 2)
para("Gợi ý được tăng dần theo bốn mức: nhắc kiến thức, chỉ vùng cần xem xét, giải thích nguyên nhân và định hướng sửa lỗi. Hệ thống không cung cấp toàn bộ lời giải ngay từ đầu, nhờ đó sinh viên vẫn phải tự suy luận và chỉnh sửa chương trình.")

doc.add_page_break()
heading("3", "Interface Wireframe")
para("Giao diện được chia thành ba khu vực: đề bài ở bên trái, trình soạn thảo và kết quả kiểm thử ở giữa, AI Tutor ở bên phải. Cách bố trí này cho phép sinh viên thực hiện toàn bộ vòng lặp học tập mà không chuyển sang màn hình khác.", after=6)
p = para(align=WD_ALIGN_PARAGRAPH.CENTER, after=3)
p.add_run().add_picture(str(IMAGE), width=Cm(12.0))
p = para(align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
font(p.add_run("Fig. 1. Wireframe của chức năng làm bài và nhận gợi ý thích ứng"), 9, italic=True)

heading("4", "Interaction Flow")
p = para(align=WD_ALIGN_PARAGRAPH.CENTER, before=4, after=7)
font(p.add_run("Đọc đề → Viết mã → Chạy/Nộp bài → Xem kiểm thử → Nhận gợi ý → Sửa mã"), 10, True)
para("Ví dụ trong wireframe minh họa sinh viên viết hàm tính tổng số chẵn nhưng sử dụng range(1, n), khiến các kiểm thử với n chẵn thất bại. AI Tutor không đưa mã sửa hoàn chỉnh mà đặt câu hỏi định hướng về biên trên của range để sinh viên tự phát hiện lỗi.")

heading("5", "Conclusion")
para("Wireframe thể hiện đầy đủ chức năng cốt lõi và mối liên hệ giữa thao tác của sinh viên với phản hồi của môi trường thực thi và AI Tutor. Bản thiết kế sẽ được sử dụng làm cơ sở cho việc xây dựng giao diện web và hiện thực hóa luồng gợi ý thích ứng trong các giai đoạn tiếp theo.")

doc.core_properties.title = "Báo cáo wireframe chức năng chính PyTutor"
doc.core_properties.author = "Tran Khanh"
doc.save(OUT)
print(OUT)
